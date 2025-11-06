import json
import logging
import os
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl import load_workbook
from app.constantes import arquivo_json, arquivo_excel
from app.exceptions import ArquivoError

def gerar_arquivo_json(nome_arquivo: str, json_conteudo: list):
    """
    Gera e salva um arquivo JSON formatado em um diretório específico,
    com logs detalhados de cada etapa do processo.

    Args:
        nome_arquivo (str): Nome base do arquivo (sem extensão).
        json_conteudo (list): Dados a serem salvos no arquivo.

    Raises:
        ArquivoError: Caso ocorra erro ao gerar o arquivo.
    """
    nome_arquivo_completo = f"{nome_arquivo}.json"
    caminho_arquivo = os.path.join(os.path.dirname(arquivo_json), nome_arquivo_completo)

    logging.debug(f"Iniciando geração do arquivo JSON: {nome_arquivo_completo}")
    logging.debug(f"Caminho completo: {caminho_arquivo}")

    try:
        with open(caminho_arquivo, "w", encoding="utf-8") as arquivo:
            json.dump(json_conteudo, arquivo, ensure_ascii=False, indent=4)

        logging.info(
            f"Arquivo '{nome_arquivo_completo}' salvo com sucesso em '{os.path.dirname(arquivo_json)}' "
            f"({len(json_conteudo)} registros)."
        )
    except Exception as e:
        logging.critical(
            f"Falha ao gerar o arquivo '{nome_arquivo_completo}': {e}"
        )
        raise ArquivoError(f"Erro ao gerar o arquivo JSON '{nome_arquivo_completo}'.") from e
    
def gerar_arquivo_excel(nome_arquivo: str, caminho_json: str = None):
    """
    Gera uma planilha Excel (.xlsx) a partir de um arquivo JSON existente,
    com formatação automática das colunas conforme o tamanho do conteúdo.

    Args:
        nome_arquivo (str): Nome base do arquivo Excel (sem extensão).
        caminho_json (str, opcional): Caminho do JSON de origem. Se não for informado,
                                      o sistema tentará detectar automaticamente o arquivo JSON
                                      com base no nome do Excel (ex: academias.json → academias.xlsx).

    Raises:
        ArquivoError: Caso ocorra erro ao gerar a planilha Excel.
    """
    try:
        caminho_origem = (
            caminho_json
            or os.path.join(os.path.dirname(arquivo_json), f"{nome_arquivo}.json")
        )

        nome_excel_completo = f"{nome_arquivo}.xlsx"
        caminho_excel = os.path.join(os.path.dirname(arquivo_excel), nome_excel_completo)

        logging.debug(f"Iniciando geração da planilha Excel: {nome_excel_completo}")
        logging.debug(f"Origem dos dados (JSON): {caminho_origem}")

        if not os.path.exists(caminho_origem):
            raise ArquivoError(f"Arquivo JSON não encontrado: {caminho_origem}")

        with open(caminho_origem, "r", encoding="utf-8") as arquivo_json_in:
            dados = json.load(arquivo_json_in)

        if not dados:
            raise ArquivoError("O arquivo JSON está vazio, nenhum dado para converter.")

        df = pd.DataFrame(dados)

        df.rename(
            columns={
                "nome": "Nome do Estabelecimento",
                "tipo": "Tipo do Estabelecimento",
                "nota": "Nota",
                "avaliacoes": "Avaliações",
                "endereco": "Endereço Completo",
            },
            inplace=True,
        )

        df.to_excel(caminho_excel, index=False, engine="openpyxl")

        wb = load_workbook(caminho_excel)
        ws = wb.active

        for cell in ws[1]:
            cell.value = str(cell.value).upper()
            cell.font = Font(bold=True, color="000000")
            cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for linha in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in linha:
                if cell.column_letter in ["C", "D"]:  # Colunas Nota e Avaliações
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        for coluna in ws.columns:
            max_len = 0
            coluna_letra = get_column_letter(coluna[0].column)
            for celula in coluna:
                try:
                    if celula.value:
                        max_len = max(max_len, len(str(celula.value)))
                except Exception:
                    pass
            ws.column_dimensions[coluna_letra].width = max_len + 2

        colunas_minimas = {"B": 30, "C": 15, "D": 20}
        for col, largura in colunas_minimas.items():
            if ws.column_dimensions[col].width < largura:
                ws.column_dimensions[col].width = largura
        
        ws.auto_filter.ref = ws.dimensions

        ws.title = nome_arquivo.capitalize()

        wb.save(caminho_excel)
        wb.close()

        logging.info(
            f"Planilha '{nome_excel_completo}' criada com sucesso em '{os.path.dirname(arquivo_excel)}' "
            f"({len(df)} registros)."
        )

    except ArquivoError as e:
        raise e
    except Exception as e:
        logging.critical(f"Erro inesperado ao gerar planilha Excel: {e}")
        raise ArquivoError(f"Falha ao gerar a planilha '{nome_arquivo}.xlsx'.") from e